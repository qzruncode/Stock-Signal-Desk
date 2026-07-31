"""Secure acquisition, extraction, and storage of text-bearing documents."""

from __future__ import annotations

from hashlib import sha256
from html import unescape
import io
import logging
import mimetypes
import os
from pathlib import Path, PurePosixPath
import re
import tempfile
from typing import Any, Iterable
from urllib.parse import unquote, urljoin, urlparse
import xml.etree.ElementTree as ET
import zipfile

import httpx

from src.agent.rss_contracts import (
    RssItemRef,
    TextChunk,
    TextDocumentResource,
)
from src.config import Config
from src.storage import DatabaseManager
from src.tools.webfetch import _validate_public_url

logger = logging.getLogger(__name__)

_DEFAULT_MAX_BYTES = 50 * 1024 * 1024
_MAX_REDIRECTS = 8
_DEFAULT_TIMEOUT = 45.0
_CHUNK_CHARS = 4_000
_CHUNK_OVERLAP = 300
_MAX_EXTRACTED_CHARS = 5_000_000
_MAX_ARCHIVE_FILES = 10_000
_MAX_ARCHIVE_UNCOMPRESSED_BYTES = 200 * 1024 * 1024

_PLAIN_TEXT_MIMES = {
    "application/atom+xml",
    "application/csv",
    "application/feed+json",
    "application/json",
    "application/ld+json",
    "application/rss+xml",
    "application/rtf",
    "application/xhtml+xml",
    "application/xml",
    "text/csv",
    "text/tab-separated-values",
}
_OFFICE_EXTENSIONS = {
    ".doc",
    ".docx",
    ".ppt",
    ".pptx",
    ".xls",
    ".xlsx",
}


class TextDocumentError(RuntimeError):
    pass


def _max_bytes() -> int:
    raw = os.getenv("AGENT_TEXT_DOCUMENT_MAX_BYTES")
    try:
        return max(1_048_576, int(raw)) if raw else _DEFAULT_MAX_BYTES
    except (TypeError, ValueError):
        return _DEFAULT_MAX_BYTES


def _resource_root() -> Path:
    configured = str(os.getenv("AGENT_TEXT_RESOURCE_DIR") or "").strip()
    if configured:
        return Path(configured).expanduser().resolve()
    database_path = Path(Config.get_instance().database_path).expanduser()
    if not database_path.is_absolute():
        database_path = Path.cwd() / database_path
    return (database_path.resolve().parent / "agent_resources").resolve()


def _safe_filename(value: str, fallback: str = "document") -> str:
    text = PurePosixPath(unquote(str(value or ""))).name
    text = re.sub(r"[\x00-\x1f/\\]+", "_", text).strip(" .")
    return (text or fallback)[:500]


def _filename_from_response(
    url: str,
    headers: httpx.Headers,
    hinted: str,
) -> str:
    disposition = str(headers.get("content-disposition") or "")
    match = re.search(
        r"filename\\*?=(?:UTF-8''|\"?)([^\";]+)",
        disposition,
        re.I,
    )
    if match:
        return _safe_filename(match.group(1))
    if hinted:
        hinted_name = _safe_filename(hinted)
        if PurePosixPath(hinted_name).suffix:
            return hinted_name
        url_extension = PurePosixPath(urlparse(url).path).suffix
        return hinted_name + url_extension if url_extension else hinted_name
    return _safe_filename(urlparse(url).path)


def _detected_mime(body: bytes, fallback: str) -> str:
    head = body[:512]
    if head.startswith(b"%PDF-"):
        return "application/pdf"
    if (
        head.startswith((b"\x89PNG\r\n\x1a\n", b"\xff\xd8\xff", b"GIF8", b"BM"))
        or head[8:12] == b"WEBP"
    ):
        return "image/unknown"
    if (
        head.startswith((b"ID3", b"OggS", b"fLaC"))
        or head[8:12] == b"WAVE"
    ):
        return "audio/unknown"
    if head[4:8] == b"ftyp" or head.startswith(b"\x1aE\xdf\xa3"):
        return "video/unknown"
    if head.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(body)) as archive:
                names = set(archive.namelist())
            if any(name.startswith("word/") for name in names):
                return (
                    "application/vnd.openxmlformats-officedocument."
                    "wordprocessingml.document"
                )
            if any(name.startswith("ppt/") for name in names):
                return (
                    "application/vnd.openxmlformats-officedocument."
                    "presentationml.presentation"
                )
            if any(name.startswith("xl/") for name in names):
                return (
                    "application/vnd.openxmlformats-officedocument."
                    "spreadsheetml.sheet"
                )
        except zipfile.BadZipFile:
            pass
    return fallback


def _validate_ooxml_archive(body: bytes) -> None:
    """Reject encrypted and decompression-bomb-shaped Office archives."""
    with zipfile.ZipFile(io.BytesIO(body)) as archive:
        members = archive.infolist()
        if len(members) > _MAX_ARCHIVE_FILES:
            raise TextDocumentError("Office 文件包含过多归档条目")
        if any(member.flag_bits & 0x1 for member in members):
            raise TextDocumentError("不支持加密的 Office 文件")
        total_size = sum(max(0, member.file_size) for member in members)
        if total_size > _MAX_ARCHIVE_UNCOMPRESSED_BYTES:
            raise TextDocumentError("Office 文件解压后超过安全上限")


def _download(
    url: str,
    *,
    title: str = "",
    declared_mime: str = "",
) -> tuple[bytes, str, str, str]:
    current_url = str(url or "").strip()
    if not current_url:
        raise TextDocumentError("文档 URL 为空")
    final_headers: httpx.Headers | None = None
    final_url = ""
    body = b""
    headers = {
        "User-Agent": (
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
            "AppleWebKit/537.36 Chrome/143.0.0.0 Safari/537.36"
        ),
        "Accept": (
            "text/*,application/pdf,application/json,application/xml,"
            "application/msword,application/vnd.ms-excel,"
            "application/vnd.ms-powerpoint,"
            "application/vnd.openxmlformats-officedocument.*;q=0.9,*/*;q=0.2"
        ),
    }
    with httpx.Client(
        follow_redirects=False,
        timeout=_DEFAULT_TIMEOUT,
    ) as client:
        for _ in range(_MAX_REDIRECTS + 1):
            _validate_public_url(current_url)
            with client.stream(
                "GET",
                current_url,
                headers=headers,
            ) as response:
                if response.status_code in {301, 302, 303, 307, 308}:
                    location = response.headers.get("location")
                    if not location:
                        raise TextDocumentError(
                            "文档重定向响应缺少 Location"
                        )
                    current_url = urljoin(
                        str(response.url),
                        location,
                    )
                    continue
                try:
                    response.raise_for_status()
                except httpx.HTTPError as exc:
                    raise TextDocumentError(
                        f"文档下载失败: {exc}"
                    ) from exc
                final_url = str(response.url)
                final_headers = response.headers
                content_length = response.headers.get(
                    "content-length"
                )
                max_bytes = _max_bytes()
                if content_length and content_length.isdigit():
                    if int(content_length) > max_bytes:
                        raise TextDocumentError(
                            f"文档超过大小上限 {max_bytes} 字节"
                        )
                chunks: list[bytes] = []
                downloaded = 0
                for chunk in response.iter_bytes():
                    downloaded += len(chunk)
                    if downloaded > max_bytes:
                        raise TextDocumentError(
                            f"文档超过大小上限 {max_bytes} 字节"
                        )
                    chunks.append(chunk)
                body = b"".join(chunks)
                break
        else:
            raise TextDocumentError(
                f"文档重定向超过 {_MAX_REDIRECTS} 次"
            )
    if final_headers is None or not final_url:
        raise TextDocumentError("文档下载没有返回最终响应")
    _validate_public_url(final_url)
    header_mime = str(final_headers.get("content-type") or "").split(
        ";",
        1,
    )[0].strip().lower()
    mime = header_mime or str(declared_mime or "").split(";", 1)[0].lower()
    filename = _filename_from_response(
        final_url,
        final_headers,
        title,
    )
    guessed_mime = mimetypes.guess_type(filename)[0] or ""
    mime = _detected_mime(
        body,
        mime or guessed_mime or "application/octet-stream",
    )
    if mime == "application/pdf":
        if not filename.lower().endswith(".pdf"):
            filename += ".pdf"
    if mime.startswith(("image/", "audio/", "video/")):
        raise TextDocumentError(f"资源是已丢弃的非文本媒体: {mime}")
    from api.v1.endpoints._rss_text import is_text_document_reference

    if not is_text_document_reference(
        url=final_url,
        mime_type=mime,
        title=filename,
    ):
        raise TextDocumentError(f"资源不是可处理的文本型文档: {mime}")
    return body, mime, filename, final_url


def _decode_text(body: bytes) -> str:
    try:
        from charset_normalizer import from_bytes

        best = from_bytes(body).best()
        if best is not None:
            return str(best)
    except Exception:
        logger.debug("Text document charset detection failed", exc_info=True)
    return body.decode("utf-8", errors="replace")


def _plain_text_from_html(value: str) -> str:
    from api.v1.endpoints._rss_text import sanitize_text_html

    sanitized = sanitize_text_html(value)
    text = re.sub(r"<br\\s*/?>", "\n", sanitized, flags=re.I)
    return unescape(re.sub(r"<[^>]+>", " ", text))


def _pdf_pages(body: bytes) -> list[tuple[int, str]]:
    from pdfminer.high_level import extract_pages
    from pdfminer.layout import LTTextContainer

    pages: list[tuple[int, str]] = []
    for page_number, layout in enumerate(
        extract_pages(io.BytesIO(body)),
        1,
    ):
        text = "\n".join(
            element.get_text()
            for element in layout
            if isinstance(element, LTTextContainer)
        )
        normalized = re.sub(r"[ \t]+\n", "\n", text).strip()
        if normalized:
            pages.append((page_number, normalized))
    return pages


def _office_text(
    body: bytes,
    *,
    extension: str,
    source_url: str,
) -> str:
    from markitdown import MarkItDown

    result = MarkItDown(enable_plugins=False).convert_stream(
        io.BytesIO(body),
        file_extension=extension,
        url=source_url,
    )
    return str(result.text_content or "").strip()


def _docx_text(body: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(body)) as archive:
        document = archive.read("word/document.xml")
    root = ET.fromstring(document)
    paragraphs: list[str] = []
    for paragraph in root.iter():
        if paragraph.tag.rsplit("}", 1)[-1] != "p":
            continue
        text = "".join(
            str(value.text or "")
            for value in paragraph.iter()
            if value.tag.rsplit("}", 1)[-1] == "t"
        ).strip()
        if text:
            paragraphs.append(text)
    return "\n\n".join(paragraphs)


def _pptx_chunks(body: bytes) -> list[dict[str, Any]]:
    with zipfile.ZipFile(io.BytesIO(body)) as archive:
        slide_names = sorted(
            (
                name
                for name in archive.namelist()
                if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)
            ),
            key=lambda name: int(
                re.search(r"slide(\d+)", name).group(1)  # type: ignore[union-attr]
            ),
        )
        chunks: list[dict[str, Any]] = []
        for slide_number, name in enumerate(slide_names, 1):
            root = ET.fromstring(archive.read(name))
            text = "\n".join(
                str(value.text or "").strip()
                for value in root.iter()
                if value.tag.rsplit("}", 1)[-1] == "t"
                and str(value.text or "").strip()
            )
            chunks.extend(
                _split_text(
                    text,
                    page=slide_number,
                    section=f"第 {slide_number} 页幻灯片",
                    start_index=len(chunks),
                )
            )
    return chunks


def _xlsx_chunks(body: bytes) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(
        io.BytesIO(body),
        read_only=True,
        data_only=True,
    )
    chunks: list[dict[str, Any]] = []
    try:
        for worksheet in workbook.worksheets:
            lines = [
                "\t".join(
                    "" if value is None else str(value)
                    for value in row
                ).rstrip()
                for row in worksheet.iter_rows(values_only=True)
            ]
            chunks.extend(
                _split_text(
                    "\n".join(line for line in lines if line),
                    section=f"工作表：{worksheet.title}",
                    start_index=len(chunks),
                )
            )
    finally:
        workbook.close()
    return chunks


def _split_text(
    text: str,
    *,
    page: int | None = None,
    section: str | None = None,
    start_index: int = 0,
) -> list[dict[str, Any]]:
    normalized = re.sub(r"\r\n?", "\n", str(text or "")).strip()
    chunks: list[dict[str, Any]] = []
    if not normalized:
        return chunks
    cursor = 0
    index = start_index
    while cursor < len(normalized):
        hard_end = min(len(normalized), cursor + _CHUNK_CHARS)
        end = hard_end
        if hard_end < len(normalized):
            boundary = max(
                normalized.rfind("\n\n", cursor, hard_end),
                normalized.rfind("。", cursor, hard_end),
                normalized.rfind("\n", cursor, hard_end),
            )
            if boundary > cursor + (_CHUNK_CHARS // 2):
                end = boundary + 1
        chunk_text = normalized[cursor:end].strip()
        if chunk_text:
            chunks.append(
                TextChunk(
                    chunk_index=index,
                    text=chunk_text,
                    content_hash=sha256(
                        chunk_text.encode("utf-8")
                    ).hexdigest(),
                    page=page,
                    section=section,
                    char_start=cursor,
                    char_end=end,
                ).model_dump(mode="json")
            )
            index += 1
        if end >= len(normalized):
            break
        cursor = max(cursor + 1, end - _CHUNK_OVERLAP)
    return chunks


def _extract(
    body: bytes,
    *,
    mime: str,
    filename: str,
    source_url: str,
) -> tuple[str, str, list[dict[str, Any]], str | None]:
    extension = PurePosixPath(filename).suffix.lower()
    if mime == "application/pdf" or extension == ".pdf":
        pages = _pdf_pages(body)
        if not pages:
            return "empty_text_layer", "pdfminer", [], None
        chunks: list[dict[str, Any]] = []
        for page, text in pages:
            chunks.extend(
                _split_text(
                    text,
                    page=page,
                    section=f"第 {page} 页",
                    start_index=len(chunks),
                )
            )
        return "extracted", "pdfminer", chunks, None
    if extension == ".docx":
        try:
            _validate_ooxml_archive(body)
            text = _docx_text(body)
        except Exception as exc:
            return "failed", "ooxml-docx", [], str(exc)
        if not text:
            return "empty_text_layer", "ooxml-docx", [], None
        return (
            "extracted",
            "ooxml-docx",
            _split_text(text[:_MAX_EXTRACTED_CHARS]),
            None,
        )
    if extension == ".pptx":
        try:
            _validate_ooxml_archive(body)
            chunks = _pptx_chunks(body)
        except Exception as exc:
            return "failed", "ooxml-pptx", [], str(exc)
        if not chunks:
            return "empty_text_layer", "ooxml-pptx", [], None
        return "extracted", "ooxml-pptx", chunks, None
    if extension == ".xlsx":
        try:
            _validate_ooxml_archive(body)
            chunks = _xlsx_chunks(body)
        except Exception as exc:
            return "failed", "openpyxl", [], str(exc)
        if not chunks:
            return "empty_text_layer", "openpyxl", [], None
        return "extracted", "openpyxl", chunks, None
    if extension in _OFFICE_EXTENSIONS:
        try:
            text = _office_text(
                body,
                extension=extension,
                source_url=source_url,
            )
        except Exception as exc:
            return "failed", "markitdown", [], str(exc)
        if not text:
            return "empty_text_layer", "markitdown", [], None
        return (
            "extracted",
            "markitdown",
            _split_text(text[:_MAX_EXTRACTED_CHARS]),
            None,
        )
    text = _decode_text(body)
    if mime in {"text/html", "application/xhtml+xml"} or extension in {
        ".html",
        ".htm",
    }:
        text = _plain_text_from_html(text)
    text = text[:_MAX_EXTRACTED_CHARS].strip()
    if not text:
        return "empty_text_layer", "text_decode", [], None
    return "extracted", "text_decode", _split_text(text), None


def _public_resource(
    value: dict[str, Any],
) -> dict[str, Any]:
    resource_id = str(value["resource_id"])
    public = {
        key: value.get(key)
        for key in (
            "resource_id",
            "filename",
            "mime_type",
            "size_bytes",
            "content_hash",
            "extraction_status",
            "extraction_method",
            "text_length",
            "chunk_count",
            "source_item_ref",
            "error",
        )
    }
    public["preview_url"] = (
        f"/api/v1/agent/resources/{resource_id}/content?disposition=inline"
    )
    public["download_url"] = (
        f"/api/v1/agent/resources/{resource_id}/content?disposition=attachment"
    )
    return TextDocumentResource.model_validate(public).model_dump(mode="json")


def materialize_text_document(
    *,
    db: DatabaseManager,
    conversation_id: str,
    run_id: str,
    url: str,
    title: str = "",
    mime_type: str = "",
    source_item_ref: dict[str, Any] | None = None,
) -> dict[str, Any]:
    body, actual_mime, filename, final_url = _download(
        url,
        title=title,
        declared_mime=mime_type,
    )
    content_hash = sha256(body).hexdigest()
    resource_id = "textdoc_" + sha256(
        f"{conversation_id}:{content_hash}".encode("utf-8")
    ).hexdigest()[:40]
    root = _resource_root()
    blob_path = root / content_hash[:2] / content_hash
    blob_path.parent.mkdir(parents=True, exist_ok=True)
    if not blob_path.exists():
        file_descriptor, temporary_name = tempfile.mkstemp(
            prefix=f".{content_hash}.",
            suffix=".tmp",
            dir=blob_path.parent,
        )
        temporary = Path(temporary_name)
        try:
            with os.fdopen(file_descriptor, "wb") as stream:
                stream.write(body)
                stream.flush()
                os.fsync(stream.fileno())
            os.replace(temporary, blob_path)
        finally:
            temporary.unlink(missing_ok=True)
    status, method, chunks, error = _extract(
        body,
        mime=actual_mime,
        filename=filename,
        source_url=final_url,
    )
    text_length = sum(len(chunk["text"]) for chunk in chunks)
    validated_item_ref = None
    if source_item_ref:
        validated_item_ref = RssItemRef.model_validate(
            source_item_ref
        ).model_dump(mode="json")
    saved = db.save_text_document(
        resource={
            "resource_id": resource_id,
            "conversation_id": conversation_id,
            "run_id": run_id,
            "source_url": final_url,
            "filename": filename,
            "mime_type": actual_mime,
            "size_bytes": len(body),
            "content_hash": content_hash,
            "blob_path": str(blob_path),
            "extraction_status": status,
            "extraction_method": method,
            "text_length": text_length,
            "source_item_ref": validated_item_ref,
            "error": error,
        },
        chunks=chunks,
    )
    return _public_resource(saved)


def materialize_text_documents(
    *,
    db: DatabaseManager,
    conversation_id: str,
    run_id: str,
    attachments: Iterable[dict[str, Any]],
    source_item_ref: dict[str, Any] | None,
) -> tuple[list[dict[str, Any]], list[str]]:
    resources: list[dict[str, Any]] = []
    errors: list[str] = []
    for attachment in attachments:
        if not isinstance(attachment, dict):
            continue
        try:
            resources.append(
                materialize_text_document(
                    db=db,
                    conversation_id=conversation_id,
                    run_id=run_id,
                    url=str(attachment.get("url") or ""),
                    title=str(attachment.get("title") or ""),
                    mime_type=str(
                        attachment.get("mime_type")
                        or attachment.get("mime")
                        or ""
                    ),
                    source_item_ref=source_item_ref,
                )
            )
        except Exception as exc:
            errors.append(
                f"{attachment.get('title') or attachment.get('url')}: {exc}"
            )
    return resources, errors


def read_text_resource(
    *,
    db: DatabaseManager,
    resource_id: str,
    offset: int = 0,
    limit: int = 50,
    page_start: int | None = None,
    page_end: int | None = None,
    query: str = "",
    conversation_id: str | None = None,
) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    document = db.get_text_document(resource_id)
    if document is None:
        raise TextDocumentError("文本资源不存在")
    if (
        conversation_id is not None
        and document.get("conversation_id") != conversation_id
    ):
        raise TextDocumentError("文本资源不属于当前会话")
    if query:
        chunks: list[dict[str, Any]] = []
        chunk_offset = 0
        while True:
            page = db.get_text_document_chunks(
                resource_id,
                offset=chunk_offset,
                limit=500,
                page_start=page_start,
                page_end=page_end,
            )
            chunks.extend(page)
            if len(page) < 500:
                break
            chunk_offset += len(page)
    else:
        chunks = db.get_text_document_chunks(
            resource_id,
            offset=offset,
            limit=limit,
            page_start=page_start,
            page_end=page_end,
        )
    query_tokens = {
        token
        for token in re.findall(
            r"[a-z0-9]+|[\u3400-\u9fff]{1,4}",
            str(query or "").lower(),
        )
        if token
    }
    if query_tokens:
        scored: list[tuple[int, dict[str, Any]]] = []
        for chunk in chunks:
            lowered = str(chunk.get("text") or "").lower()
            score = sum(
                lowered.count(token) for token in query_tokens
            )
            if score:
                scored.append((score, chunk))
        scored.sort(
            key=lambda value: (
                -value[0],
                int(value[1].get("chunk_index") or 0),
            )
        )
        chunks = [value for _score, value in scored[:limit]]
    return _public_resource(document), chunks[:limit]


def delete_conversation_document_blobs(
    db: DatabaseManager,
    conversation_id: str,
) -> None:
    blob_paths = db.delete_text_documents(conversation_id)
    for value in blob_paths:
        if not value or db.text_document_blob_is_referenced(value):
            continue
        path = Path(value)
        try:
            root = _resource_root()
            resolved = path.resolve()
            if root not in resolved.parents:
                logger.error(
                    "Refusing to delete text resource outside root: %s",
                    resolved,
                )
                continue
            resolved.unlink(missing_ok=True)
        except Exception:
            logger.warning(
                "Failed to delete text resource blob %s",
                value,
                exc_info=True,
            )


__all__ = [
    "TextDocumentError",
    "delete_conversation_document_blobs",
    "materialize_text_document",
    "materialize_text_documents",
    "read_text_resource",
]
